import html
import io
import json
import os

import pandas as pd
import streamlit as st
from langchain_core.prompts import PromptTemplate
from langchain_groq import ChatGroq

try:
    from langchain_classic.output_parsers import (
        OutputFixingParser,
        ResponseSchema,
        StructuredOutputParser,
    )
except ImportError:
    from langchain.output_parsers import (
        OutputFixingParser,
        ResponseSchema,
        StructuredOutputParser,
    )

MAX_CHARS = 24000
DEFAULT_MODEL = "openai/gpt-oss-20b"

# ----------------------------------------------------------------------------
# 1. Output schema + prompt 
# ----------------------------------------------------------------------------
response_schemas = [
    ResponseSchema(
        name="full_name",
        description="The candidate's full name as a string",
    ),
    ResponseSchema(
        name="email",
        description="The candidate's email address as a string",
    ),
    ResponseSchema(
        name="education",
        description=(
            "A list of education entries. Each entry is a JSON object with keys: "
            "'degree' (string), 'institution' (string), 'year' (integer)."
        ),
    ),
    ResponseSchema(
        name="skills",
        description="A list of strings representing the candidate's skills",
    ),
    ResponseSchema(
        name="experience",
        description=(
            "A list of experience entries. Each entry is a JSON object with keys: "
            "'role' (string), 'company' (string), 'years' (string, e.g. '2020-2023' or '2023-Present')."
        ),
    ),
]

output_parser = StructuredOutputParser.from_response_schemas(response_schemas)

TEMPLATE = """You are an expert HR assistant that extracts structured information from resume text.

Read the resume snippet below carefully and extract the candidate's information.
Return ONLY the data in the exact format specified below, with no extra commentary.

Resume Text:
{resume_text}

{format_instructions}
"""

prompt = PromptTemplate(
    template=TEMPLATE,
    input_variables=["resume_text"],
    partial_variables={"format_instructions": output_parser.get_format_instructions()},
)


# ----------------------------------------------------------------------------
# 2. Helpers
# ----------------------------------------------------------------------------
def extract_text_from_upload(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if name.endswith(".docx"):
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


def _as_list(value):
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except Exception:
            return [value] if value.strip() else []
    if value is None:
        return []
    return value if isinstance(value, list) else [value]


def parse_resume(resume_text: str, api_key: str, model: str, temperature: float) -> dict:
    llm = ChatGroq(model=model, temperature=temperature, api_key=api_key)
    fixing_parser = OutputFixingParser.from_llm(parser=output_parser, llm=llm)

    raw_output = llm.invoke(prompt.format_prompt(resume_text=resume_text).to_string())
    try:
        result = output_parser.parse(raw_output.content)
    except Exception:
        result = fixing_parser.parse(raw_output.content)

    for key in ("education", "skills", "experience"):
        result[key] = _as_list(result.get(key))
    return result


def render_skills(skills):
    chips = "".join(
        f"<span style='display:inline-block;margin:3px 6px 3px 0;padding:4px 12px;"
        f"border-radius:14px;background:rgba(99,102,241,.15);border:1px solid rgba(99,102,241,.4);"
        f"font-size:0.9rem'>{html.escape(str(s))}</span>"
        for s in skills
    )
    st.markdown(chips or "_No skills found._", unsafe_allow_html=True)


# ----------------------------------------------------------------------------
# 3. UI
# ----------------------------------------------------------------------------
st.set_page_config(page_title="HR Candidate Profile Parser", page_icon="📄", layout="wide")
st.title("📄 HR Candidate Profile Parser")
st.caption("Upload a CV (PDF / DOCX / TXT) or paste its text, and get a structured candidate profile.")

env_key = os.environ.get("GROQ_API_KEY", "")

with st.sidebar:
    st.header("Settings")
    if env_key:
        api_key = env_key
        st.success("Groq API key loaded from the server.")
    else:
        api_key = st.text_input("Groq API key", type="password", help="Get a free key from console.groq.com")
    model = st.text_input("Model", value=DEFAULT_MODEL)
    temperature = st.slider("Temperature", 0.0, 1.0, 0.0, 0.1)

tab_upload, tab_paste = st.tabs(["Upload CV", "Paste text"])
with tab_upload:
    uploaded = st.file_uploader("CV file", type=["pdf", "docx", "txt"])
with tab_paste:
    pasted = st.text_area("Resume text", height=250)

if st.button("Parse profile", type="primary"):
    try:
        if uploaded is not None:
            resume_text = extract_text_from_upload(uploaded)
        else:
            resume_text = pasted
        resume_text = (resume_text or "").strip()

        if not api_key:
            st.error("Please enter your Groq API key in the sidebar.")
        elif not resume_text:
            st.error("No text found. Upload a text-based CV or paste the resume text. (Scanned PDFs are not supported.)")
        else:
            if len(resume_text) > MAX_CHARS:
                st.warning(f"The CV is long, only the first {MAX_CHARS:,} characters will be used.")
                resume_text = resume_text[:MAX_CHARS]
            with st.spinner("Analyzing the CV..."):
                st.session_state["result"] = parse_resume(resume_text, api_key, model.strip(), temperature)
                st.session_state["resume_text"] = resume_text
    except Exception as exc:
        st.session_state.pop("result", None)
        st.error(f"Something went wrong: {exc}")

result = st.session_state.get("result")
if result:
    st.divider()
    col1, col2 = st.columns(2)
    col1.metric("Full name", result.get("full_name") or "-")
    col2.metric("Email", result.get("email") or "-")

    st.subheader("🎓 Education")
    if result["education"]:
        st.dataframe(pd.DataFrame(result["education"]), hide_index=True)
    else:
        st.write("No education entries found.")

    st.subheader("🛠️ Skills")
    render_skills(result["skills"])

    st.subheader("💼 Experience")
    if result["experience"]:
        st.dataframe(pd.DataFrame(result["experience"]), hide_index=True)
    else:
        st.write("No experience entries found.")

    st.divider()
    json_str = json.dumps(result, indent=2, ensure_ascii=False)
    st.download_button("⬇️ Download JSON", json_str, file_name="candidate_profile.json", mime="application/json")
    with st.expander("Raw JSON"):
        st.code(json_str, language="json")
    with st.expander("Extracted CV text"):
        st.text(st.session_state.get("resume_text", ""))
